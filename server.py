import os
import json
import asyncio
import threading
import concurrent.futures
import pandas as pd
from fastapi import FastAPI, BackgroundTasks, HTTPException
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sse_starlette.sse import EventSourceResponse

from scrapers.jobspy_scraper import get_jobspy_jobs
from scrapers.himalayas_scraper import get_himalayas_jobs
from scrapers.remotive_scraper import get_remotive_jobs
from scrapers.remoteok_scraper import get_remoteok_jobs
from scrapers.weworkremotely_scraper import get_weworkremotely_jobs
from scrapers.nodesk_python_scraper import get_nodesk_python_jobs
from evaluator import evaluate_job_single
from database import save_job_to_db, get_saved_jobs_from_db, is_mongo_connected
from log_manager import log_queue, job_queue, pending_job_queue, log

app = FastAPI()

# Absolute path resolution
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")

# Mount static files for the frontend using absolute path
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

global_evaluated_jobs = []
global_pending_jobs = []
global_source_stats = {}
is_running = False

# Lock for thread-safe access to global lists
list_lock = threading.Lock()

class VerifyJobRequest(BaseModel):
    url: str
    action: str  # "approve" or "reject"
    match_score: int
    match_reason: str
    estimated_pay: str
    location: str

def run_pipeline():
    global is_running, global_evaluated_jobs, global_pending_jobs, global_source_stats
    is_running = True
    global_evaluated_jobs = []
    global_pending_jobs = []
    global_source_stats = {}
    
    log("[SYSTEM] --- Phase 1: Multi-Site Data Extraction Started ---")
    
    scraped_batches = []
    
    # Define tasks for concurrent scraping execution
    scraper_tasks = [
        ("JobSpy (LinkedIn/Indeed/Glassdoor/ZipRecruiter)", get_jobspy_jobs),
        ("Himalayas API", get_himalayas_jobs),
        ("Remotive API", get_remotive_jobs),
        ("Remote OK API", get_remoteok_jobs),
        ("WeWorkRemotely RSS", get_weworkremotely_jobs),
        ("Python.org & NoDesk RSS", get_nodesk_python_jobs)
    ]
    
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as executor:
        future_map = {executor.submit(func): name for name, func in scraper_tasks}
        for future in concurrent.futures.as_completed(future_map):
            name = future_map[future]
            try:
                jobs = future.result()
                if jobs:
                    scraped_batches.extend(jobs)
                    count = len(jobs)
                    global_source_stats[name] = count
                    log(f"[SCRAPER] Source '{name}' completed successfully -> Found {count} raw jobs")
                else:
                    global_source_stats[name] = 0
                    log(f"[SCRAPER] Source '{name}' returned 0 jobs")
            except Exception as e:
                global_source_stats[name] = 0
                log(f"[ERROR] Error executing scraper {name}: {e}")
                
    log(f"[SYSTEM] Total raw jobs collected across all sources: {len(scraped_batches)}")
    
    # Deduplicate by URL and Title/Company fingerprint
    unique_jobs = {}
    for job in scraped_batches:
        url = job.get("url", "").strip()
        fingerprint = f"{job.get('title','').strip().lower()}:{job.get('company','').strip().lower()}"
        if url and url not in unique_jobs and fingerprint not in unique_jobs:
            unique_jobs[url] = job
            unique_jobs[fingerprint] = job

    # Filter out fingerprint keys to leave list of unique job dicts
    jobs_list = [j for k, j in unique_jobs.items() if k.startswith("http") or k.startswith("https")]
    log(f"[SYSTEM] Total unique jobs after deduplication: {len(jobs_list)}")
    
    if not jobs_list:
        log("[SYSTEM] No jobs found across scraped sources. Pipeline complete.")
        is_running = False
        return
        
    log("[SYSTEM] --- Phase 2: AI Evaluation & Fit Scoring Engine Started ---")
    log(f"[EVALUATOR] Evaluating {len(jobs_list)} unique jobs in parallel against candidate CV...")
    
    def evaluate_and_track(job):
        res = evaluate_job_single(job)
        if res:
            with list_lock:
                if res.get('requires_manual_verification', False):
                    global_pending_jobs.append(res)
                elif res.get('match_score', 0) >= 40 and res.get('is_pakistan_eligible', False):
                    global_evaluated_jobs.append(res)
                    # Automatically save to MongoDB if connected
                    save_job_to_db(res)
        return res

    with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
        futures = [executor.submit(evaluate_and_track, job) for job in jobs_list]
        concurrent.futures.wait(futures)
    
    log("[SYSTEM] --- Phase 3: Extraction & Matching Pipeline Completed ---")
    log(f"[SUMMARY] Pipeline finished! Found {len(global_evaluated_jobs)} matched jobs in Right Window.")
    log(f"[SUMMARY] Found {len(global_pending_jobs)} jobs pending manual verification.")
    if is_mongo_connected:
        log(f"[DATABASE] Automatically persisted {len(global_evaluated_jobs)} matched jobs to MongoDB.")
    log("DONE")
    is_running = False

@app.get("/")
def serve_index():
    index_file = os.path.join(STATIC_DIR, "index.html")
    with open(index_file, "r", encoding="utf-8") as f:
        return HTMLResponse(content=f.read())

@app.post("/api/start")
def start_scraping(background_tasks: BackgroundTasks):
    global is_running
    if is_running:
        return {"status": "Already running"}
    background_tasks.add_task(run_pipeline)
    return {"status": "Started"}

@app.get("/api/stream")
async def sse_stream():
    async def event_generator():
        while True:
            while not log_queue.empty():
                msg = log_queue.get()
                yield {"event": "log", "data": json.dumps({"message": msg, "sources": global_source_stats, "mongo_connected": is_mongo_connected})}
                
            while not job_queue.empty():
                job = job_queue.get()
                try:
                    job_json = json.dumps(job, default=str)
                    yield {"event": "job", "data": job_json}
                except Exception as e:
                    log(f"Error serializing job: {e}")
            
            while not pending_job_queue.empty():
                job = pending_job_queue.get()
                try:
                    job_json = json.dumps(job, default=str)
                    yield {"event": "pending_job", "data": job_json}
                except Exception as e:
                    log(f"Error serializing pending job: {e}")
            
            if not is_running and log_queue.empty() and job_queue.empty() and pending_job_queue.empty():
                yield {"event": "done", "data": json.dumps({"status": "complete", "sources": global_source_stats, "mongo_connected": is_mongo_connected})}
                break
                
            await asyncio.sleep(0.5)
            
    return EventSourceResponse(event_generator())

@app.get("/api/jobs/saved")
def get_db_jobs():
    """
    Returns saved jobs stored in MongoDB.
    """
    if not is_mongo_connected:
        return {"mongo_connected": False, "jobs": global_evaluated_jobs}
    saved_jobs = get_saved_jobs_from_db()
    return {"mongo_connected": True, "jobs": saved_jobs}

@app.post("/api/verify-job")
def verify_job(req: VerifyJobRequest):
    global global_evaluated_jobs, global_pending_jobs
    
    target_job = None
    with list_lock:
        for job in global_pending_jobs:
            if job.get("url") == req.url:
                target_job = job
                break
                
        if not target_job:
            raise HTTPException(status_code=404, detail="Job not found in pending verification list")
            
        global_pending_jobs.remove(target_job)
        
        if req.action == "approve":
            target_job["match_score"] = req.match_score
            target_job["match_reason"] = f"[Manually Verified] {req.match_reason}"
            target_job["estimated_pay"] = req.estimated_pay
            target_job["location"] = req.location
            target_job["is_pakistan_eligible"] = True
            target_job["requires_manual_verification"] = False
            
            global_evaluated_jobs.append(target_job)
            save_job_to_db(target_job)
            return {"status": "approved", "job": target_job}
        else:
            return {"status": "rejected"}

@app.get("/api/download")
def download_csv():
    if not global_evaluated_jobs:
        return {"error": "No jobs available for download"}
        
    df = pd.DataFrame(global_evaluated_jobs)
    cols = ["match_score", "title", "company", "estimated_pay", "location", "url", "source", "match_reason", "description"]
    cols = [c for c in cols if c in df.columns]
    df = df[cols]
    df = df.sort_values(by="match_score", ascending=False)
    
    output_file = os.path.join(BASE_DIR, "remote_junior_jobs.csv")
    df.to_csv(output_file, index=False)
    return FileResponse(path=output_file, filename="remote_junior_jobs.csv", media_type="text/csv")

@app.get("/api/download/docx")
def download_docx():
    if not global_evaluated_jobs:
        return {"error": "No jobs available for download"}
        
    output_file = os.path.join(BASE_DIR, "remote_junior_jobs.docx")
    try:
        create_docx_report(global_evaluated_jobs, output_file)
        return FileResponse(path=output_file, filename="remote_junior_jobs.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        return {"error": f"Failed to generate DOCX: {e}"}

def create_docx_report(jobs, filepath):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = docx.Document()
    
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("AI MATCHED REMOTE JOBS REPORT")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(59, 130, 246)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"Generated on {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')} | Remote Jobs (Pakistan Eligible)")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(148, 163, 184)
    
    desc_p = doc.add_paragraph()
    desc_p.add_run("\nThe following list displays remote job opportunities matching your CV. All listed positions allow remote work from anywhere (Pakistan-eligible) and match the junior technical profile.\n")
    
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Light Shading Accent 1'
    
    hdr_cells = table.rows[0].cells
    headers = ["Score", "Job Title & Company", "Estimated Pay", "Location Restriction", "Match Reason"]
    widths = [Inches(0.8), Inches(2.2), Inches(1.2), Inches(1.2), Inches(2.6)]
    
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = header
        hdr_cells[idx].width = widths[idx]
        for paragraph in hdr_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.name = 'Arial'
                
    sorted_jobs = sorted(jobs, key=lambda x: x.get('match_score', 0), reverse=True)
    
    for job in sorted_jobs:
        row_cells = table.add_row().cells
        score = job.get('match_score', 0)
        row_cells[0].text = f"{score}%"
        row_cells[0].width = widths[0]
        
        title_p = row_cells[1].paragraphs[0]
        run_title = title_p.add_run(job.get('title', ''))
        run_title.font.bold = True
        run_title.font.size = Pt(9.5)
        run_title.font.name = 'Arial'
        
        comp_p = row_cells[1].add_paragraph()
        run_comp = comp_p.add_run(f"{job.get('company', '')} ({job.get('source', '')})")
        run_comp.font.italic = True
        run_comp.font.size = Pt(9)
        run_comp.font.name = 'Arial'
        
        url = job.get('url', '')
        if url:
            link_p = row_cells[1].add_paragraph()
            run_url = link_p.add_run("Apply: " + url)
            run_url.font.size = Pt(8)
            run_url.font.color.rgb = RGBColor(59, 130, 246)
                
        row_cells[1].width = widths[1]
        
        row_cells[2].text = job.get('estimated_pay', 'Not Disclosed')
        row_cells[2].width = widths[2]
        
        row_cells[3].text = job.get('location', 'Remote')
        row_cells[3].width = widths[3]
        
        row_cells[4].text = job.get('match_reason', '')
        row_cells[4].width = widths[4]
        
        for idx in [0, 2, 3, 4]:
            for p in row_cells[idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Arial'
                    
    doc.save(filepath)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=True)
